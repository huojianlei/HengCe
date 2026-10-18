# -*- coding: utf-8 -*-
"""v2 轨道（九表 + 财会2019-6号三表）导入器测试。

覆盖：
    1) 正例：生成器产物（预算基础 / 实时经营）整本导入通过，计数与分录数正确；
    2) 拒收：表名被改 → 表序列门；
    3) 拒收：静态标签被改 → 标签门；
    4) 拒收：数值位写成文本 → 数值位门；
    5) 拒收：利润表营业利润被改 → 三表恒等式门；
    6) 拒收：分录单笔借贷不平 → 分录门。

测试数据来自生成器实跑（零外部依赖：没有生成器产物就跳过）。
"""
import io
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

BEN_MU = Path(__file__).resolve().parent.parent               # ingest_v11
sys.path.insert(0, str(BEN_MU))                               # 本包优先

from excel_ingest_v2 import V2_SHEETS, import_v2              # noqa: E402
from xlsx_reader import ExcelError                            # noqa: E402

CHAN = BEN_MU.parent / '演示生成器v2' / '输出'                 # 生成器产物目录
ZHENG = CHAN / '演示-预算基础-20261001.xlsx'                   # 正例样本（预算基础）
SHI = CHAN / '演示-实时经营-20261001.xlsx'                     # 正例样本（实时经营）


def _gai(lu, gai_dong):
    """复制一本并改动若干格：gai_dong = [(表, 格位, 新值)]，返回字节。"""
    wb = load_workbook(lu)
    for biao, ref, zhi in gai_dong:
        wb[biao][ref] = zhi
    liu = io.BytesIO()
    wb.save(liu)
    return liu.getvalue()


@unittest.skipUnless(ZHENG.exists(), '缺生成器产物（先跑 演示生成器v2/shengcheng.py）')
class V2DaoruTest(unittest.TestCase):
    """v2 轨道导入器：正例 + 五条拒收路径。"""

    def test_zhengli_yusuan(self):
        """预算基础册：整本通过，分录 531 笔（12 月 × 43 + 年末结转 15）。"""
        r = import_v2(ZHENG.read_bytes(), ZHENG.name)
        self.assertTrue(r['ok'])
        self.assertEqual(r['contract_version'], 'v2')
        self.assertEqual(r['category'], '预算基础')
        self.assertEqual(tuple(r['sheets']), V2_SHEETS)
        self.assertEqual(len(r['fenlu']), 531)
        self.assertGreater(r['checks'], 10000)

    def test_zhengli_shishi(self):
        """实时经营册：6 个月，分录 273 笔（6 × 43 + 15）。"""
        r = import_v2(SHI.read_bytes(), SHI.name)
        self.assertEqual(r['category'], '实时经营')
        self.assertEqual(len(r['fenlu']), 273)
        self.assertEqual(len(r['tanxiao']['台账']), 8)

    def test_jushou_biaoming(self):
        """表名被改 → 表序列门拒收。"""
        wb = load_workbook(ZHENG)
        wb['分录'].title = '分录表'
        liu = io.BytesIO()
        wb.save(liu)
        with self.assertRaises(ExcelError):
            import_v2(liu.getvalue(), 'x.xlsx')

    def test_jushou_biaoqian(self):
        """静态标签被改（分录表头）→ 标签门拒收。"""
        yu = _gai(ZHENG, [('分录', 'B1', '摘要（改）')])
        with self.assertRaises(ExcelError):
            import_v2(yu, 'x.xlsx')

    def test_jushou_shuzhi(self):
        """数值位写成文本 → 数值位门拒收。"""
        yu = _gai(ZHENG, [('摊销明细', 'B3', '两千一')])
        with self.assertRaises(ExcelError):
            import_v2(yu, 'x.xlsx')

    def test_jushou_hengdengshi(self):
        """利润表营业利润被改 → 三表恒等式门拒收。"""
        yu = _gai(ZHENG, [('利润表', 'B23', 1)])
        with self.assertRaises(ExcelError):
            import_v2(yu, 'x.xlsx')

    def test_jushou_fenlu(self):
        """分录单笔借贷不平 → 分录门拒收。"""
        yu = _gai(ZHENG, [('分录', 'F2', 1)])
        with self.assertRaises(ExcelError):
            import_v2(yu, 'x.xlsx')


if __name__ == '__main__':
    unittest.main(verbosity=2)
