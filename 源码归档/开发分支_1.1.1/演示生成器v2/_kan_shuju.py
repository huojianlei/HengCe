# -*- coding: utf-8 -*-
"""九表体检（开发期查看）：打印九表关键数字与空格分布。

三表用 _kan_san.py 看。
用法：python -B _kan_shuju.py [产物路径]
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

from banshi import SHEETS

MU = Path(__file__).resolve().parent
JIU = ('税金及附加', '销售费用', '管理费用',
       '费用表月度汇总', '摊销明细', '分录')


def s(v):
    return f'{v:,.2f}' if isinstance(v, (int, float)) else '—'


def kan(lu):
    wb = load_workbook(lu)
    print(f'== {Path(lu).name} ==')
    for b in ('销售费用', '管理费用'):
        ws = wb[b]
        print(f'-- {b}：年合计 {s(ws["O34"].value)}（含税 {s(ws["O36"].value)}）')
    ws = wb['税金及附加']
    print(f'-- 税金及附加：附加率 {ws["C31"].value} 年税 {s(ws["AC127"].value)}')
    tx = wb['摊销明细']
    print(f'-- 摊销明细：原值 {s(tx["C31"].value)} 月数 {tx["C32"].value} '
          f'月摊销 {s(tx["C33"].value)} 矩阵合计 {s(tx["AB142"].value)}')
    fl = wb['分录']
    bi = sum(1 for r in range(124, 461) if fl[f'O{r}'].value)
    print(f'-- 分录：{bi} 笔；借 {s(fl["S463"].value)} 贷 {s(fl["T464"].value)} '
          f'差 {fl["U465"].value}')
    print('-- 空/满分布（已用网格内）：')
    for mb in JIU:
        ws = wb[mb]
        shu = wen = kong = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    kong += 1
                elif isinstance(c.value, (int, float)):
                    shu += 1
                else:
                    wen += 1
        print(f'   {mb:<10} 数字 {shu:>5}  文本 {wen:>4}  空 {kong:>5}  '
              f'网格 {ws.max_row}×{ws.max_column}')
    wb.close()


if __name__ == '__main__':
    p = Path(sys.argv[1]) if len(sys.argv) > 1 else (
        MU / '输出' / '演示-预算基础-20260923.xlsx')
    kan(p)
