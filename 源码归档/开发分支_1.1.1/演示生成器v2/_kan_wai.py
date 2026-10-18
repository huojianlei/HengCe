# -*- coding: utf-8 -*-
"""外部两册体检（开发期查看）：打印批次信息/基准指标/事件台账/科目影响。

用法：python -B _kan_wai.py [产物路径]
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

MU = Path(__file__).resolve().parent


def kan(lu):
    wb = load_workbook(lu)
    print(f'== {Path(lu).name} ==')
    for ws in wb.worksheets:
        print(f'-- {ws.title}（{ws.max_row}×{ws.max_column}）')
        jian = 0
        for row in ws.iter_rows():
            zhi = [(c.coordinate, c.value) for c in row if c.value is not None]
            if not zhi:
                continue
            print('   ', ' | '.join(f'{k}={v}' for k, v in zhi)[:190])
            jian += 1
            if jian >= 8:
                print('    …')
                break
    wb.close()


if __name__ == '__main__':
    for m in ('行业对标', '经营变更'):
        p = Path(sys.argv[1]) if len(sys.argv) > 1 else (MU / '输出' / f'演示-{m}-20260923.xlsx')
        kan(p)
        print()
